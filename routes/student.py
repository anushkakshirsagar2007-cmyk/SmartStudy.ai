import os
import uuid
import json
import base64
import io
from datetime import datetime, timedelta
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, current_app, abort, send_from_directory
from flask_login import login_required, current_user
from models import User, Note, Task, Document, Quiz, Question, QuizAttempt, FlashcardSet, Flashcard, ChatMessage, StudySession, FREE_UPLOAD_LIMIT
from extensions import db
from sqlalchemy import func

student_bp = Blueprint('student', __name__)

ALLOWED_EXTENSIONS = {'pdf', 'txt', 'docx', 'md', 'png', 'jpg', 'jpeg', 'webp'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def transcribe_scanned_pdf(file_path, subject, filename):
    """OCR locally when available, then use the configured vision model on bounded pages."""
    from routes.api import ai_error_response, require_openai

    try:
        import pypdfium2
        pdf = pypdfium2.PdfDocument(file_path)
        try:
            page_limit = min(len(pdf), current_app.config['SCANNED_PDF_PAGE_LIMIT'])
            images = [pdf[index].render(scale=1.5).to_pil() for index in range(page_limit)]
        finally:
            pdf.close()
    except ImportError:
        return '[Scanned PDF detected, but PDF rendering is unavailable. Install pypdfium2 to process scanned PDFs.]'
    except Exception:
        current_app.logger.exception('Failed to render scanned PDF %s', filename)
        return '[Scanned PDF detected, but its pages could not be rendered.]'

    try:
        import pytesseract
        ocr_text = '\n'.join(pytesseract.image_to_string(image) for image in images).strip()
        if len(ocr_text) >= 50:
            return ocr_text
    except (ImportError, OSError):
        pass
    except Exception:
        current_app.logger.info('Optional local OCR could not process %s', filename, exc_info=True)

    client, model, error = require_openai()
    if error:
        return '[Scanned PDF detected, but transcription requires OPENAI_API_KEY or working local pytesseract OCR.]'

    content = [{'type': 'text', 'text': f'Transcribe all readable text, equations, and labels from these first {page_limit} pages. Preserve page order. Subject: {subject}.'}]
    for image in images:
        image_buffer = io.BytesIO()
        image.save(image_buffer, format='JPEG', quality=85)
        encoded = base64.b64encode(image_buffer.getvalue()).decode('ascii')
        content.append({
            'type': 'image_url',
            'image_url': {'url': f'data:image/jpeg;base64,{encoded}'}
        })
    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{'role': 'user', 'content': content}],
            max_tokens=3000
        )
        transcription = (response.choices[0].message.content or '').strip()
        return transcription or '[Scanned PDF transcription returned no text.]'
    except Exception as exc:
        result, status = ai_error_response(exc, 'Scanned PDF transcription')
        error = result.get_json().get('error', 'Scanned PDF transcription failed.')
        current_app.logger.warning('Vision PDF transcription failed with status %s', status)
        return f'[Scanned PDF transcription failed: {error}]'


@student_bp.app_context_processor
def inject_now():
    return {'now': datetime.utcnow()}


# ─── Dashboard ─────────────────────────────────────────────────────────────────

@student_bp.route('/dashboard')
@login_required
def dashboard():
    if current_user.is_admin():
        return redirect(url_for('admin.dashboard'))
    
    # Stats
    total_notes = current_user.notes.count()
    total_tasks = current_user.tasks.count()
    completed_tasks = current_user.tasks.filter_by(status='completed').count()
    total_quizzes = current_user.quizzes.count()
    total_docs = current_user.documents.count()
    total_flashcard_sets = current_user.flashcard_sets.count()
    
    # Recent quiz performance
    recent_attempts = current_user.quiz_attempts.order_by(QuizAttempt.completed_at.desc()).limit(5).all()
    avg_score = 0
    if recent_attempts:
        avg_score = sum(a.score for a in recent_attempts) / len(recent_attempts)
    
    # Recent notes
    recent_notes = current_user.notes.order_by(Note.updated_at.desc()).limit(4).all()
    
    # Upcoming tasks
    upcoming_tasks = current_user.tasks.filter(
        Task.status != 'completed',
        Task.due_date >= datetime.utcnow()
    ).order_by(Task.due_date.asc()).limit(5).all()
    
    # Study streak (simplified - count days with activity in last 7 days)
    week_ago = datetime.utcnow() - timedelta(days=7)
    recent_activity = current_user.study_sessions.filter(
        StudySession.created_at >= week_ago
    ).all()
    active_days = len(set(s.created_at.date() for s in recent_activity))
    
    # Recent documents
    recent_docs = current_user.documents.order_by(Document.created_at.desc()).limit(3).all()
    
    stats = {
        'total_notes': total_notes,
        'total_tasks': total_tasks,
        'completed_tasks': completed_tasks,
        'task_completion_rate': round((completed_tasks / total_tasks * 100) if total_tasks else 0, 1),
        'total_quizzes': total_quizzes,
        'total_docs': total_docs,
        'total_flashcard_sets': total_flashcard_sets,
        'avg_quiz_score': round(avg_score, 1),
        'active_days': active_days,
        'total_quiz_attempts': current_user.quiz_attempts.count()
    }
    
    return render_template('student/dashboard.html',
                           stats=stats,
                           recent_notes=recent_notes,
                           upcoming_tasks=upcoming_tasks,
                           recent_attempts=recent_attempts,
                           recent_docs=recent_docs)


# ─── Notes ─────────────────────────────────────────────────────────────────────

@student_bp.route('/notes')
@login_required
def notes():
    search = request.args.get('search', '')
    subject_filter = request.args.get('subject', '')
    bookmarked = request.args.get('bookmarked', '')
    
    query = current_user.notes
    
    if search:
        query = query.filter(
            (Note.title.ilike(f'%{search}%')) | (Note.content.ilike(f'%{search}%'))
        )
    if subject_filter:
        query = query.filter_by(subject=subject_filter)
    if bookmarked == '1':
        query = query.filter_by(is_bookmarked=True)
    
    notes_list = query.order_by(Note.updated_at.desc()).all()
    subjects = db.session.query(Note.subject).filter_by(user_id=current_user.id).distinct().all()
    subjects = [s[0] for s in subjects if s[0]]
    
    return render_template('student/notes.html', notes=notes_list, subjects=subjects,
                           search=search, subject_filter=subject_filter, bookmarked=bookmarked)


@student_bp.route('/notes/create', methods=['GET', 'POST'])
@login_required
def create_note():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '').strip()
        subject = request.form.get('subject', 'General').strip()
        tags = request.form.get('tags', '').strip()
        color = request.form.get('color', '#7C3AED')
        
        if not title:
            flash('Note title is required.', 'error')
            return redirect(url_for('student.create_note'))
        
        note = Note(
            user_id=current_user.id,
            title=title,
            content=content,
            subject=subject,
            tags=tags,
            color=color
        )
        db.session.add(note)
        
        # Log study session
        session_log = StudySession(user_id=current_user.id, activity_type='notes', subject=subject)
        db.session.add(session_log)
        db.session.commit()
        
        flash('Note created successfully!', 'success')
        return redirect(url_for('student.notes'))
    
    return render_template('student/note_editor.html', note=None, mode='create')


@student_bp.route('/notes/<int:note_id>/edit', methods=['GET', 'POST'])
@login_required
def edit_note(note_id):
    note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
    
    if request.method == 'POST':
        note.title = request.form.get('title', '').strip() or note.title
        note.content = request.form.get('content', '').strip()
        note.subject = request.form.get('subject', 'General').strip()
        note.tags = request.form.get('tags', '').strip()
        note.color = request.form.get('color', note.color)
        note.updated_at = datetime.utcnow()
        db.session.commit()
        
        flash('Note updated successfully!', 'success')
        return redirect(url_for('student.notes'))
    
    return render_template('student/note_editor.html', note=note, mode='edit')


@student_bp.route('/notes/<int:note_id>/delete', methods=['POST'])
@login_required
def delete_note(note_id):
    note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
    db.session.delete(note)
    db.session.commit()
    flash('Note deleted.', 'info')
    return redirect(url_for('student.notes'))


@student_bp.route('/notes/<int:note_id>/toggle-bookmark', methods=['POST'])
@login_required
def toggle_bookmark(note_id):
    note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
    note.is_bookmarked = not note.is_bookmarked
    db.session.commit()
    return jsonify({'bookmarked': note.is_bookmarked})


# ─── Tasks / Planner ───────────────────────────────────────────────────────────

@student_bp.route('/planner')
@login_required
def planner():
    filter_status = request.args.get('status', '')
    filter_priority = request.args.get('priority', '')
    filter_subject = request.args.get('subject', '')
    
    query = current_user.tasks
    if filter_status:
        query = query.filter_by(status=filter_status)
    if filter_priority:
        query = query.filter_by(priority=filter_priority)
    if filter_subject:
        query = query.filter_by(subject=filter_subject)
    
    tasks = query.order_by(Task.due_date.asc().nullsfirst(), Task.created_at.desc()).all()
    subjects = db.session.query(Task.subject).filter_by(user_id=current_user.id).distinct().all()
    subjects = [s[0] for s in subjects if s[0]]
    
    # Group by status
    pending = [t for t in tasks if t.status == 'pending']
    in_progress = [t for t in tasks if t.status == 'in_progress']
    completed = [t for t in tasks if t.status == 'completed']
    
    return render_template('student/planner.html', 
                           tasks=tasks, pending=pending,
                           in_progress=in_progress, completed=completed,
                           subjects=subjects,
                           filter_status=filter_status,
                           filter_priority=filter_priority,
                           filter_subject=filter_subject)


@student_bp.route('/planner/task/create', methods=['POST'])
@login_required
def create_task():
    title = request.form.get('title', '').strip()
    description = request.form.get('description', '').strip()
    subject = request.form.get('subject', 'General').strip()
    priority = request.form.get('priority', 'medium')
    due_date_str = request.form.get('due_date', '')
    
    if not title:
        flash('Task title is required.', 'error')
        return redirect(url_for('student.planner'))
    
    due_date = None
    if due_date_str:
        try:
            due_date = datetime.strptime(due_date_str, '%Y-%m-%dT%H:%M')
        except ValueError:
            try:
                due_date = datetime.strptime(due_date_str, '%Y-%m-%d')
            except ValueError:
                pass
    
    task = Task(
        user_id=current_user.id,
        title=title,
        description=description,
        subject=subject,
        priority=priority,
        due_date=due_date
    )
    db.session.add(task)
    db.session.commit()
    
    flash('Task added successfully!', 'success')
    return redirect(url_for('student.planner'))


@student_bp.route('/planner/task/<int:task_id>/update-status', methods=['POST'])
@login_required
def update_task_status(task_id):
    task = Task.query.filter_by(id=task_id, user_id=current_user.id).first_or_404()
    new_status = request.form.get('status', task.status)
    
    valid_statuses = ['pending', 'in_progress', 'completed']
    if new_status in valid_statuses:
        task.status = new_status
        if new_status == 'completed':
            task.completed_at = datetime.utcnow()
        else:
            task.completed_at = None
        db.session.commit()
        flash('Task status updated.', 'success')
    
    return redirect(url_for('student.planner'))


@student_bp.route('/planner/task/<int:task_id>/delete', methods=['POST'])
@login_required
def delete_task(task_id):
    task = Task.query.filter_by(id=task_id, user_id=current_user.id).first_or_404()
    db.session.delete(task)
    db.session.commit()
    flash('Task deleted.', 'info')
    return redirect(url_for('student.planner'))


# ─── Documents / AI Workspace ──────────────────────────────────────────────────

@student_bp.route('/workspace')
@login_required
def workspace():
    docs = current_user.documents.order_by(Document.created_at.desc()).all()
    return render_template('student/workspace.html', documents=docs)


@student_bp.route('/workspace/upload', methods=['POST'])
@login_required
def upload_document():
    # ─── UPLOAD LIMIT CHECK ────────────────────────────────────────────────────
    if not current_user.can_upload:
        flash(f'You have used all {FREE_UPLOAD_LIMIT} free uploads. Upgrade to Pro for unlimited uploads!', 'error')
        return redirect(url_for('payment.pricing'))
    
    if 'file' not in request.files:
        flash('No file selected.', 'error')
        return redirect(url_for('student.workspace'))
    
    file = request.files['file']
    subject = request.form.get('subject', 'General').strip()
    
    if not file or file.filename == '':
        flash('No file selected.', 'error')
        return redirect(url_for('student.workspace'))
    
    if not allowed_file(file.filename):
        flash('File type not supported. Please upload PDF, TXT, DOCX, or MD files.', 'error')
        return redirect(url_for('student.workspace'))
    
    # Secure storage
    ext = file.filename.rsplit('.', 1)[1].lower()
    stored_name = f"{current_user.id}_{uuid.uuid4().hex}.{ext}"
    upload_path = current_app.config['UPLOAD_FOLDER']
    
    # Create user-specific subdirectory
    user_upload_dir = os.path.join(upload_path, str(current_user.id))
    os.makedirs(user_upload_dir, exist_ok=True)
    
    file_path = os.path.join(user_upload_dir, stored_name)
    file.save(file_path)
    file_size = os.path.getsize(file_path)
    
    # Extract text content with improved handling
    content_text = ''
    page_count = 0
    
    try:
        if ext == 'pdf':
            import pypdf
            with open(file_path, 'rb') as f:
                reader = pypdf.PdfReader(f)
                page_count = len(reader.pages)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        content_text += page_text + '\n'
            content_text = content_text.strip()
            if len(content_text) < 50:
                content_text = transcribe_scanned_pdf(file_path, subject, file.filename)
        elif ext == 'txt' or ext == 'md':
            # Try multiple encodings
            for encoding in ['utf-8', 'latin-1', 'cp1252', 'ascii']:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content_text = f.read()
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue
            page_count = 1
        elif ext in ('docx', 'doc'):
            from docx import Document as DocxDocument
            doc_obj = DocxDocument(file_path)
            paragraphs = [p.text for p in doc_obj.paragraphs if p.text.strip()]
            content_text = '\n'.join(paragraphs)
            page_count = max(1, len(content_text) // 3000)  # Estimate pages
        elif ext in ('png', 'jpg', 'jpeg', 'webp'):
            page_count = 1
            # ─── Multimodal Vision AI Analysis ─────────────────────────────────
            from routes.api import require_openai
            client, model, _ = require_openai()
            if client:
                import base64
                with open(file_path, 'rb') as img_f:
                    b64_data = base64.b64encode(img_f.read()).decode('utf-8')
                mime = 'image/jpeg' if ext in ('jpg', 'jpeg') else f'image/{ext}'
                try:
                    vision_resp = client.chat.completions.create(
                        model=model,
                        messages=[{
                            'role': 'user',
                            'content': [
                                {
                                    'type': 'text',
                                    'text': (
                                        f"You are an expert academic tutor for {subject}.\n"
                                        "Please analyze and comprehensively explain this study image/diagram in detail:\n"
                                        "1. Detailed Transcription: Transcribe all visible text, formulas, math equations, or labels.\n"
                                        "2. Diagram & Visual Analysis: Explain any diagrams, charts, graphs, flowcharts, anatomical structures, or circuits shown.\n"
                                        "3. Comprehensive Concept Explanation: Provide an in-depth academic explanation of the theory and principles illustrated in this image so a student can study directly from it.\n"
                                        "4. Key Study Takeaways: Highlight the critical points, exam tips, and definitions."
                                    )
                                },
                                {
                                    'type': 'image_url',
                                    'image_url': {
                                        'url': f'data:{mime};base64,{b64_data}'
                                    }
                                }
                            ]
                        }],
                        max_tokens=2500
                    )
                    content_text = vision_resp.choices[0].message.content
                except Exception as ve:
                    content_text = f"Image study document: {file.filename}\n[Vision analysis note: {str(ve)}]"
            else:
                content_text = f"Image study document: {file.filename}\n[Uploaded image. Add OPENAI_API_KEY to generate full vision breakdown.]"
    except Exception as e:
        content_text = f'[Could not extract text: {str(e)}]'
    
    word_count = len(content_text.split()) if content_text else 0
    
    doc = Document(
        user_id=current_user.id,
        original_name=file.filename,
        stored_name=stored_name,
        file_type=ext,
        file_size=file_size,
        content_text=content_text[:50000],  # Limit stored text
        subject=subject,
        is_processed=bool(content_text and not content_text.startswith('[')),
        page_count=page_count,
        word_count=word_count
    )
    db.session.add(doc)
    
    # ─── INCREMENT FREE UPLOAD COUNTER ─────────────────────────────────────────
    if not current_user.is_pro:
        current_user.free_uploads_used += 1
    
    session_log = StudySession(user_id=current_user.id, activity_type='document', subject=subject)
    db.session.add(session_log)
    db.session.commit()
    
    if content_text.startswith('[Scanned PDF'):
        flash(content_text[1:-1], 'warning')
    else:
        flash(f'Document "{file.filename}" uploaded and processed successfully!', 'success')
    return redirect(url_for('student.document_view', doc_id=doc.id))


@student_bp.route('/workspace/document/<int:doc_id>')
@login_required
def document_view(doc_id):
    doc = Document.query.filter_by(id=doc_id, user_id=current_user.id).first_or_404()
    related_notes = Note.query.filter_by(user_id=current_user.id, source_document_id=doc_id).all()
    related_quizzes = Quiz.query.filter_by(user_id=current_user.id, source_document_id=doc_id).all()
    related_flashcards = FlashcardSet.query.filter_by(user_id=current_user.id, source_document_id=doc_id).all()
    has_openai = bool(current_app.config.get('OPENAI_API_KEY') or os.environ.get('OPENAI_API_KEY'))
    
    return render_template('student/document_view.html', 
                           doc=doc,
                           related_notes=related_notes,
                           related_quizzes=related_quizzes,
                           related_flashcards=related_flashcards,
                           has_openai=has_openai)


@student_bp.route('/workspace/file/<int:doc_id>')
@login_required
def serve_document_file(doc_id):
    """Serve uploaded document, PDF, or image file securely for inline preview."""
    doc = Document.query.filter_by(id=doc_id, user_id=current_user.id).first_or_404()
    user_upload_dir = os.path.join(current_app.config['UPLOAD_FOLDER'], str(current_user.id))
    mimetype = None
    if doc.file_type == 'pdf':
        mimetype = 'application/pdf'
    elif doc.file_type in ('png', 'webp'):
        mimetype = f'image/{doc.file_type}'
    elif doc.file_type in ('jpg', 'jpeg'):
        mimetype = 'image/jpeg'
    elif doc.file_type in ('txt', 'md'):
        mimetype = 'text/plain; charset=utf-8'
    return send_from_directory(user_upload_dir, doc.stored_name, mimetype=mimetype, as_attachment=False)


@student_bp.route('/workspace/document/<int:doc_id>/delete', methods=['POST'])
@login_required
def delete_document(doc_id):
    doc = Document.query.filter_by(id=doc_id, user_id=current_user.id).first_or_404()
    
    # Remove file
    user_upload_dir = os.path.join(current_app.config['UPLOAD_FOLDER'], str(current_user.id))
    file_path = os.path.join(user_upload_dir, doc.stored_name)
    if os.path.exists(file_path):
        os.remove(file_path)
    
    db.session.delete(doc)
    db.session.commit()
    
    flash('Document deleted.', 'info')
    return redirect(url_for('student.workspace'))


# ─── Quizzes ───────────────────────────────────────────────────────────────────

@student_bp.route('/quizzes')
@login_required
def quizzes():
    quizzes_list = current_user.quizzes.order_by(Quiz.created_at.desc()).all()
    attempts = current_user.quiz_attempts.order_by(QuizAttempt.completed_at.desc()).limit(10).all()
    return render_template('student/quizzes.html', quizzes=quizzes_list, attempts=attempts)


@student_bp.route('/quizzes/<int:quiz_id>/take')
@login_required
def take_quiz(quiz_id):
    quiz = Quiz.query.filter_by(id=quiz_id, user_id=current_user.id).first_or_404()
    questions = quiz.questions.order_by(Question.order_num).all()
    
    if not questions:
        flash('This quiz has no questions yet.', 'warning')
        return redirect(url_for('student.quizzes'))
    
    session_log = StudySession(user_id=current_user.id, activity_type='quiz', 
                                activity_id=quiz_id, subject=quiz.subject)
    db.session.add(session_log)
    db.session.commit()
    
    return render_template('student/take_quiz.html', quiz=quiz, questions=questions)


@student_bp.route('/quizzes/<int:quiz_id>/submit', methods=['POST'])
@login_required
def submit_quiz(quiz_id):
    quiz = Quiz.query.filter_by(id=quiz_id, user_id=current_user.id).first_or_404()
    questions = quiz.questions.order_by(Question.order_num).all()
    
    user_answers = {}
    correct = 0
    
    for q in questions:
        answer = request.form.get(f'q_{q.id}', '')
        user_answers[str(q.id)] = answer
        if answer.upper() == q.correct_answer.upper():
            correct += 1
    
    total = len(questions)
    score = (correct / total * 100) if total else 0
    time_taken = int(request.form.get('time_taken', 0))
    
    attempt = QuizAttempt(
        user_id=current_user.id,
        quiz_id=quiz_id,
        score=score,
        total_questions=total,
        correct_answers=correct,
        time_taken=time_taken,
        answers_json=json.dumps(user_answers)
    )
    db.session.add(attempt)
    db.session.commit()
    
    return redirect(url_for('student.quiz_result', attempt_id=attempt.id))


@student_bp.route('/quizzes/result/<int:attempt_id>')
@login_required
def quiz_result(attempt_id):
    attempt = QuizAttempt.query.filter_by(id=attempt_id, user_id=current_user.id).first_or_404()
    quiz = attempt.quiz
    questions = quiz.questions.order_by(Question.order_num).all()
    user_answers = json.loads(attempt.answers_json)
    
    return render_template('student/quiz_result.html',
                           attempt=attempt, quiz=quiz,
                           questions=questions, user_answers=user_answers)


@student_bp.route('/quizzes/<int:quiz_id>/delete', methods=['POST'])
@login_required
def delete_quiz(quiz_id):
    quiz = Quiz.query.filter_by(id=quiz_id, user_id=current_user.id).first_or_404()
    db.session.delete(quiz)
    db.session.commit()
    flash('Quiz deleted.', 'info')
    return redirect(url_for('student.quizzes'))


# ─── Flashcards ────────────────────────────────────────────────────────────────

@student_bp.route('/flashcards')
@login_required
def flashcards():
    sets = current_user.flashcard_sets.order_by(FlashcardSet.created_at.desc()).all()
    return render_template('student/flashcards.html', sets=sets)


@student_bp.route('/flashcards/set/<int:set_id>')
@login_required
def flashcard_set(set_id):
    fset = FlashcardSet.query.filter_by(id=set_id, user_id=current_user.id).first_or_404()
    cards = fset.flashcards.order_by(Flashcard.order_num).all()
    
    session_log = StudySession(user_id=current_user.id, activity_type='flashcard',
                                activity_id=set_id, subject=fset.subject)
    db.session.add(session_log)
    db.session.commit()
    
    return render_template('student/flashcard_study.html', fset=fset, cards=cards)


@student_bp.route('/flashcards/card/<int:card_id>/toggle-mastered', methods=['POST'])
@login_required
def toggle_card_mastered(card_id):
    card = Flashcard.query.join(FlashcardSet).filter(
        Flashcard.id == card_id,
        FlashcardSet.user_id == current_user.id
    ).first_or_404()
    
    card.is_mastered = not card.is_mastered
    card.times_reviewed += 1
    db.session.commit()
    
    return jsonify({'mastered': card.is_mastered, 'times_reviewed': card.times_reviewed})


@student_bp.route('/flashcards/set/<int:set_id>/delete', methods=['POST'])
@login_required
def delete_flashcard_set(set_id):
    fset = FlashcardSet.query.filter_by(id=set_id, user_id=current_user.id).first_or_404()
    db.session.delete(fset)
    db.session.commit()
    flash('Flashcard set deleted.', 'info')
    return redirect(url_for('student.flashcards'))


# ─── Progress / Analytics ──────────────────────────────────────────────────────

@student_bp.route('/progress')
@login_required
def progress():
    # All-time stats
    total_quiz_attempts = current_user.quiz_attempts.count()
    all_attempts = current_user.quiz_attempts.all()
    avg_score = sum(a.score for a in all_attempts) / len(all_attempts) if all_attempts else 0
    
    # Tasks by status
    task_stats = db.session.query(Task.status, func.count(Task.id)).filter_by(
        user_id=current_user.id
    ).group_by(Task.status).all()
    task_data = {status: count for status, count in task_stats}
    
    # Subjects breakdown
    subject_notes = db.session.query(Note.subject, func.count(Note.id)).filter_by(
        user_id=current_user.id
    ).group_by(Note.subject).all()
    
    # Weekly activity (last 7 days)
    week_ago = datetime.utcnow() - timedelta(days=7)
    weekly_sessions = current_user.study_sessions.filter(
        StudySession.created_at >= week_ago
    ).all()
    
    daily_counts = {}
    for i in range(7):
        day = (datetime.utcnow() - timedelta(days=i)).strftime('%Y-%m-%d')
        daily_counts[day] = 0
    
    for s in weekly_sessions:
        day = s.created_at.strftime('%Y-%m-%d')
        if day in daily_counts:
            daily_counts[day] += 1
    
    # Quiz score history
    quiz_history = current_user.quiz_attempts.order_by(
        QuizAttempt.completed_at.asc()
    ).limit(20).all()
    
    # Flashcard progress
    total_cards = 0
    mastered_cards = 0
    for fset in current_user.flashcard_sets.all():
        cards = fset.flashcards.all()
        total_cards += len(cards)
        mastered_cards += sum(1 for c in cards if c.is_mastered)
    
    return render_template('student/progress.html',
                           total_quiz_attempts=total_quiz_attempts,
                           avg_score=round(avg_score, 1),
                           task_data=task_data,
                           subject_notes=subject_notes,
                           daily_counts=daily_counts,
                           quiz_history=quiz_history,
                           total_cards=total_cards,
                           mastered_cards=mastered_cards)


# ─── Chat ──────────────────────────────────────────────────────────────────────

@student_bp.route('/chat')
@login_required
def chat():
    docs = current_user.documents.filter_by(is_processed=True).order_by(Document.created_at.desc()).all()
    recent_messages = current_user.chat_messages.order_by(
        ChatMessage.created_at.desc()
    ).limit(50).all()
    recent_messages = list(reversed(recent_messages))
    has_openai = bool(current_app.config.get('OPENAI_API_KEY') or os.environ.get('OPENAI_API_KEY'))
    
    return render_template('student/chat.html', 
                           docs=docs, 
                           messages=recent_messages,
                           has_openai=has_openai)


# ─── Profile / Settings ────────────────────────────────────────────────────────

@student_bp.route('/profile', methods=['GET', 'POST'])
@login_required
def profile():
    if request.method == 'POST':
        action = request.form.get('action', 'update_profile')
        
        if action == 'update_profile':
            current_user.full_name = request.form.get('full_name', '').strip() or current_user.full_name
            current_user.bio = request.form.get('bio', '').strip()
            current_user.preferred_subject = request.form.get('preferred_subject', '').strip()
            
            try:
                current_user.study_goal_hours = float(request.form.get('study_goal_hours', 2.0))
            except ValueError:
                pass
            
            avatar_color = request.form.get('avatar_color', '')
            if avatar_color and avatar_color.startswith('#'):
                current_user.avatar_color = avatar_color
            
            db.session.commit()
            flash('Profile updated successfully!', 'success')
        
        elif action == 'change_password':
            from routes.auth import validate_password
            current_password = request.form.get('current_password', '')
            new_password = request.form.get('new_password', '')
            confirm_password = request.form.get('confirm_password', '')
            
            if not current_user.check_password(current_password):
                flash('Current password is incorrect.', 'error')
            elif new_password != confirm_password:
                flash('New passwords do not match.', 'error')
            else:
                errors = validate_password(new_password)
                if errors:
                    for e in errors:
                        flash(e, 'error')
                else:
                    current_user.set_password(new_password)
                    db.session.commit()
                    flash('Password changed successfully!', 'success')
        
        return redirect(url_for('student.profile'))
    
    return render_template('student/profile.html')
